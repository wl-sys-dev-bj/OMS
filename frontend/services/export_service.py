import csv
import io
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def export_queryset_to_csv(queryset, headers, row_extractor, filename="export.csv"):
    """Export queryset to CSV response."""
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    writer = csv.writer(response)
    writer.writerow(headers)
    for obj in queryset:
        writer.writerow(row_extractor(obj))

    return response


def export_queryset_to_excel(queryset, headers, row_extractor, filename="export.xlsx", sheet_title="Report"):
    """Export queryset to styled Excel (.xlsx) response using openpyxl."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    # Styling definitions
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="173C32", end_color="173C32", fill_type="solid")
    regular_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style="thin", color="DCE4DE"),
        right=Side(style="thin", color="DCE4DE"),
        top=Side(style="thin", color="DCE4DE"),
        bottom=Side(style="thin", color="DCE4DE"),
    )

    # Write headers
    ws.append(headers)
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    ws.row_dimensions[1].height = 26

    # Write rows
    for row_num, obj in enumerate(queryset, 2):
        row_values = row_extractor(obj)
        ws.append(row_values)
        for col_num, _ in enumerate(row_values, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    # Auto-adjust column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_employees_csv(queryset):
    headers = [
        "Employee ID", "Full Name", "Phone", "Email", "Department",
        "Designation", "Employment Type", "Joining Date", "Salary (AED)",
        "Status", "Work Location", "Client Name", "Passport No", "Citizenship No"
    ]
    def extract(emp):
        return [
            emp.employee_id,
            emp.full_name,
            emp.phone,
            emp.email,
            emp.department.name if emp.department else "",
            emp.designation.name if emp.designation else (emp.job_title or ""),
            emp.get_employment_type_display(),
            str(emp.joining_date or ""),
            str(emp.salary or ""),
            emp.get_status_display(),
            emp.work_location,
            emp.client_name,
            emp.passport_number,
            emp.citizenship_number,
        ]
    return export_queryset_to_csv(queryset, headers, extract, "employees_directory.csv")


def export_employees_excel(queryset):
    headers = [
        "Employee ID", "Full Name", "Phone", "Email", "Department",
        "Designation", "Employment Type", "Joining Date", "Salary (AED)",
        "Status", "Work Location", "Client Name", "Passport No", "Citizenship No"
    ]
    def extract(emp):
        return [
            emp.employee_id,
            emp.full_name,
            emp.phone,
            emp.email,
            emp.department.name if emp.department else "",
            emp.designation.name if emp.designation else (emp.job_title or ""),
            emp.get_employment_type_display(),
            str(emp.joining_date or ""),
            float(emp.salary) if emp.salary else 0.0,
            emp.get_status_display(),
            emp.work_location,
            emp.client_name,
            emp.passport_number,
            emp.citizenship_number,
        ]
    return export_queryset_to_excel(queryset, headers, extract, "employees_directory.xlsx", "Employees")


def export_expenses_csv(queryset):
    headers = [
        "Reference No", "Bill Date", "Vendor / Payee", "Employee",
        "Category", "Amount", "Currency", "Status", "Payment Method",
        "Invoice No", "Due Date", "Paid Date", "Description"
    ]
    def extract(exp):
        return [
            exp.reference_number,
            str(exp.bill_date),
            exp.vendor_name,
            exp.employee.full_name if exp.employee else "",
            exp.category,
            str(exp.amount),
            exp.currency,
            exp.get_status_display(),
            exp.get_payment_method_display(),
            exp.invoice_number,
            str(exp.due_date or ""),
            str(exp.paid_date or ""),
            exp.description,
        ]
    return export_queryset_to_csv(queryset, headers, extract, "expense_register.csv")


def export_expenses_excel(queryset):
    headers = [
        "Reference No", "Bill Date", "Vendor / Payee", "Employee",
        "Category", "Amount", "Currency", "Status", "Payment Method",
        "Invoice No", "Due Date", "Paid Date", "Description"
    ]
    def extract(exp):
        return [
            exp.reference_number,
            str(exp.bill_date),
            exp.vendor_name,
            exp.employee.full_name if exp.employee else "",
            exp.category,
            float(exp.amount) if exp.amount else 0.0,
            exp.currency,
            exp.get_status_display(),
            exp.get_payment_method_display(),
            exp.invoice_number,
            str(exp.due_date or ""),
            str(exp.paid_date or ""),
            exp.description,
        ]
    return export_queryset_to_excel(queryset, headers, extract, "expense_register.xlsx", "Expenses")
